"""services/timesheet_sheet.py — 工時 Google Sheet 的讀取（整本 xlsx export）。

拉取 runner（services/timesheet_puller）與歷史匯入腳本（scripts/import_timesheets.py）
共用這一份：分頁名、欄位、壞列判定、日期字串格式各寫一份就會漂，而 hash 用的是
日期**字串**——兩條路格式不同就會在交接處重複入庫（docs/TIMESHEET_IMPORT_PLAN.md D9）。

試算表是公開連結：`https://docs.google.com/spreadsheets/d/<id>/export?format=xlsx` 不用
授權就拿得到整本，「總表（勿動）」已是兩個輸入分頁的聯集、值已算好（助理那頁的
IMPORTRANGE 值也在裡面）—— 所以不需要 Apps Script 推，主控端定時拉就好。
零新依賴：HTTP 用 urllib、xlsx 用 openpyxl（生產 python_embed 已裝 3.1.5）。
"""
from __future__ import annotations

import datetime as dt
import io
import urllib.request

DATA_SHEET = "總表（勿動）"
STATUS_SHEET = "專案狀態(勿動)"
_UA = "OMG-timesheet/1.0"


def export_url(sheet_id: str) -> str:
    return f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"


def fetch_xlsx(sheet_id: str, timeout: int = 60) -> bytes:
    """抓整本 xlsx（公開連結；Google 會 302 幾次，urllib 自己跟）。"""
    req = urllib.request.Request(export_url(sheet_id), headers={"User-Agent": _UA})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read()


def load_workbook(data: bytes):
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)


def to_num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def date_str(v) -> str:
    """xlsx 的 datetime → 'yyyy/MM/dd'（同 Apps Script 的 formatDate）；已是字串就原樣 strip。"""
    if isinstance(v, (dt.datetime, dt.date)):
        return f"{v.year:04d}/{v.month:02d}/{v.day:02d}"
    return str(v or "").strip()


def read_rows(wb) -> tuple:
    """「總表」→ (可送的列, 壞列)。壞列＝日期解析不了或時數不是數字 —— 收進報告，不送。
    可送列的形狀就是 ingest 的 TimesheetRow：{date, staff, project, task, hours}。"""
    ws = wb[DATA_SHEET]
    good, bad = [], []
    for i, r in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
        r = list(r) + [None] * 8
        d, who, proj, task, hrs = r[:5]
        if d is None and not who and not proj and not task:
            continue
        h = to_num(hrs)
        ds = date_str(d)
        if h is None or not ds or len(ds) < 8:
            bad.append({"row": i, "date": ds, "staff": who, "project": proj, "hours": hrs})
            continue
        good.append({"date": ds, "staff": str(who or "").strip(),
                     "project": str(proj or "").strip(),
                     "task": str(task or "").strip(), "hours": h})
    return good, bad


def read_budgets(wb) -> list:
    """「專案狀態」：預算＝剩餘＋實際（規劃 §1.1 驗過）；≤0＝沒設，不送。"""
    out = []
    for r in wb[STATUS_SHEET].iter_rows(min_row=2, values_only=True):
        r = list(r) + [None] * 12
        menu = (r[1] or "")
        if not menu:
            continue
        b = (to_num(r[9]) or 0) + (to_num(r[10]) or 0)
        if b > 0:
            out.append({"sheet_name": str(menu).strip(), "budget_hours": round(b, 1)})
    return out
