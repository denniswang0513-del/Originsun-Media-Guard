# -*- coding: utf-8 -*-
"""主控端定時拉工時 Google Sheet（services/timesheet_puller）—— owner 2026-09-03
「定期向 Google Sheet 拉資料就好了嗎？那個連結是公開連結」。

釘的規則：讀表只有一份（腳本與 runner 同吃）、寫入只有一條（router 的 ingest 端點與
runner 都走 services.timesheet_ingest）、只在 master 跑、設定驗證、拉不到不炸。
"""
import datetime as dt
import io

import pytest
from tests.unit._srcscan import code_only, func_body, repo_src, timesheets_src


@pytest.fixture
def pull_store(monkeypatch):
    """settings.json 的替身：只有 timesheet.pull 那一塊（puller 與 SettingsBlock 都讀它）。"""
    import services.timesheet_settings as tss
    store = {"timesheet": {"pull": {}}}
    monkeypatch.setattr(tss, "load_settings", lambda: store)
    monkeypatch.setattr(tss, "save_settings", lambda s: store.update(s))
    return store


def _wb(rows, budgets=()):
    """用 openpyxl 在記憶體造一本跟真表同形的 xlsx：總表列 4 起 A–E；專案狀態 B＝名、J＝剩餘、K＝實際。"""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "總表（勿動）"
    for _ in range(3):
        ws.append([None])
    for r in rows:
        ws.append(list(r))
    st = wb.create_sheet("專案狀態(勿動)")
    st.append(["#", "案名"])
    for name, remain, actual in budgets:
        st.append([None, name] + [None] * 7 + [remain, actual])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_read_rows_keeps_apps_script_date_format_and_reports_bad_rows():
    from services.timesheet_sheet import load_workbook, read_budgets, read_rows
    data = _wb([
        (dt.datetime(2026, 9, 1), "王", "客戶_案子", "剪接", 2.5),
        (dt.datetime(2026, 9, 2), "王", "行政庶務", "", None),      # 沒時數 → 壞列
        ("", "李", "客戶_案子", "", 1),                               # 沒日期 → 壞列
        (None, None, None, None, None),                            # 全空 → 跳過
    ], budgets=[("客戶_案子", 10, 30), ("沒預算", 0, 0)])
    wb = load_workbook(data)
    good, bad = read_rows(wb)
    assert good == [{"date": "2026/09/01", "staff": "王", "project": "客戶_案子", "task": "剪接", "hours": 2.5}]
    assert [b["row"] for b in bad] == [5, 6]
    assert read_budgets(wb) == [{"sheet_name": "客戶_案子", "budget_hours": 40.0}]


def test_one_read_one_write_path():
    """讀表：腳本與 runner 都 import services.timesheet_sheet；寫入：router 的 ingest 端點
    與 runner 都呼叫 services.timesheet_ingest.ingest（router 不再自己 add(Timesheet)）。"""
    router = code_only(timesheets_src())
    body = func_body(router, "async def ingest_rows(")
    assert "await ingest(session, req.rows, req.source)" in body
    assert "Timesheet(" not in body and "row_hash" not in body.replace("X-Timesheet-Token", "")
    puller = code_only(repo_src("services/timesheet_puller.py"))
    assert "from services.timesheet_ingest import ingest" in puller
    assert "from services.timesheet_sheet import" in puller
    assert 'ingest(session, chunk, "sheet", ctx)' in puller, "runner 落庫的 source 要是 sheet"
    # 查表建一次給 20 批共用，不每批重建
    assert puller.count("ingest_context(") == 1


def test_scheduler_only_runs_on_master():
    """全機隊共用同一顆 DB —— 沒 gate 每台 agent 都會各拉一份。"""
    src = code_only(repo_src("services/timesheet_puller.py"))
    loop = func_body(src, "async def _scheduler_loop(")
    assert "is_master_machine()" in loop
    assert loop.index("is_master_machine()") < loop.index("while True")
    assert "timesheet_puller.start_scheduler_task()" in repo_src("main.py")


def test_settings_validation_and_defaults(pull_store):
    import services.timesheet_puller as tp
    assert tp.DEFAULT_CRON == "0 9 * * 6"        # owner 2026-09-03：每週六一次
    cfg = tp.get_pull_settings()
    assert cfg == {"enabled": False, "sheet_id": "", "cron": tp.DEFAULT_CRON,
                   "last_run_at": 0.0, "last_summary": "", "running": False}
    cfg = tp.update_pull_settings({"enabled": True, "sheet_id": " abc ", "cron": "*/30 * * * *"})
    assert cfg["enabled"] is True and cfg["sheet_id"] == "abc" and cfg["cron"] == "*/30 * * * *"
    with pytest.raises(ValueError):
        tp.update_pull_settings({"cron": "not a cron"})
    # 端點把 ValueError 變 422，不是 500
    ep = func_body(code_only(timesheets_src()), "async def put_pull(")
    assert "except ValueError" in ep and "422" in ep


async def test_run_pull_is_disabled_by_default_and_needs_sheet_id(pull_store):
    import services.timesheet_puller as tp
    assert (await tp.run_pull()) == {"status": "disabled"}
    r = await tp.run_pull(force=True)
    assert r["status"] == "error" and "sheet_id" in r["message"]


async def test_run_pull_survives_a_fetch_failure(monkeypatch, pull_store):
    """Google 那邊掛了／網路斷了：記到 last_summary、回 error，不讓排程 loop 炸掉。"""
    import services.timesheet_puller as tp
    store = pull_store
    store["timesheet"]["pull"].update({"enabled": True, "sheet_id": "x"})

    def boom(_sid):
        raise OSError("HTTP 503")
    monkeypatch.setattr(tp, "fetch_xlsx", boom)
    r = await tp.run_pull()
    assert r["status"] == "error" and "503" in r["message"]
    assert "失敗" in store["timesheet"]["pull"]["last_summary"]
    assert tp.get_pull_settings()["running"] is False
