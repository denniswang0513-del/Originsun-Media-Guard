# -*- coding: utf-8 -*-
"""工時 Google Sheet 歷史匯入（docs/TIMESHEET_IMPORT_PLAN.md Phase B／C）。

    # dry-run：只出報告，不寫任何東西
    .venv/Scripts/python.exe scripts/import_timesheets.py --xlsx timesheet.xlsx --budget
    # 真的匯（dev 8001）；--prod 改打生產 8000
    .venv/Scripts/python.exe scripts/import_timesheets.py --xlsx timesheet.xlsx --budget --apply [--prod]

xlsx 來源：整本試算表 `export?format=xlsx`（「總表（勿動）」已是兩個輸入分頁的
聯集、值已算好）。**走 HTTP 打 /timesheets/ingest**，不直接碰 DB —— 同 row_hash、
同手填優先去重，重跑第二次 0 新增（冪等）。

🔴 日期一律轉成 `yyyy/MM/dd` 字串再送：Apps Script 是 `Utilities.formatDate(d,
'Asia/Taipei','yyyy/MM/dd')`，hash 用的是日期**字串**，兩條路格式不同就會在交接處
重複入庫（規劃 D9）。
"""
from __future__ import annotations

import argparse
import collections
import datetime as dt
import io
import json
import os
import sys
import urllib.error
import urllib.request

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, REPO)

from core.hr_logic import INTERNAL_BUCKETS, sheet_project_key  # noqa: E402

DATA_SHEET = "總表（勿動）"
STATUS_SHEET = "專案狀態(勿動)"
BATCH = 200


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _date_str(v) -> str:
    """xlsx 的 datetime → 'yyyy/MM/dd'（同 Apps Script）；已是字串就原樣 strip。"""
    if isinstance(v, (dt.datetime, dt.date)):
        return f"{v.year:04d}/{v.month:02d}/{v.day:02d}"
    return str(v or "").strip()


def read_rows(xlsx: str) -> tuple:
    """回 (可送的列, 壞列)。壞列＝日期解析不了或時數不是數字 —— 收進報告，不送。"""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[DATA_SHEET]
    good, bad = [], []
    for i, r in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
        r = list(r) + [None] * 8
        d, who, proj, task, hrs = r[:5]
        if d is None and not who and not proj and not task:
            continue
        h = _num(hrs)
        ds = _date_str(d)
        if h is None or not ds or len(ds) < 8:
            bad.append({"row": i, "date": ds, "staff": who, "project": proj, "hours": hrs})
            continue
        good.append({"date": ds, "staff": str(who or "").strip(),
                     "project": str(proj or "").strip(),
                     "task": str(task or "").strip(), "hours": h})
    return good, bad


def read_budgets(xlsx: str) -> list:
    """「專案狀態」：預算＝剩餘＋實際（規劃 §1.1 驗過）；≤0＝沒設，不送。"""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    out = []
    for r in wb[STATUS_SHEET].iter_rows(min_row=2, values_only=True):
        r = list(r) + [None] * 12
        menu = (r[1] or "")
        if not menu:
            continue
        b = (_num(r[9]) or 0) + (_num(r[10]) or 0)
        if b > 0:
            out.append({"sheet_name": str(menu).strip(), "budget_hours": round(b, 1)})
    return out


class Api:
    def __init__(self, base: str):
        self.base = base
        from core.auth import create_token
        self.jwt = create_token({"sub": "admin", "username": "import_timesheets",
                                 "access_level": 3})
        self.token = self.call("GET", "/api/v1/timesheets/ingest_token")["token"]

    def call(self, method: str, path: str, body=None, headers=None):
        data = json.dumps(body).encode() if body is not None else None
        req = urllib.request.Request(self.base + path, data=data, method=method)
        req.add_header("Authorization", "Bearer " + self.jwt)
        req.add_header("Content-Type", "application/json")
        for k, v in (headers or {}).items():
            req.add_header(k, v)
        try:
            with urllib.request.urlopen(req, timeout=120) as x:
                return json.loads(x.read().decode() or "{}")
        except urllib.error.HTTPError as e:
            raise SystemExit(f"{method} {path} → HTTP {e.code}: {e.read().decode()[:300]}")


def _dry_run(good: list, hours, P, prod: bool = False) -> None:
    """不寫：在本機用同一支 resolver 預估對映分層（讀 DB，不動 DB）。
    `prod=True` 讀生產庫（只讀 —— 對映預估要看的是生產的私帳 413 案，dev 只有幾案）。"""
    import asyncio

    import db.session as _ds
    if prod:
        _orig = _ds.get_database_url
        _ds.get_database_url = lambda: _orig().replace("mediaguard_dev", "mediaguard")
    from core.hr_logic import resolve_project, suggest_projects
    from routers.api_timesheets import _project_lookup

    async def _pre():
        await _ds.init_db()
        async with _ds.get_session_factory()() as s:
            return await _project_lookup(s)
    pmap, by_name, by_key = asyncio.run(_pre())
    tiers = collections.defaultdict(list)
    for p in sorted(hours, key=lambda x: -hours[x]):
        pid, why = resolve_project(p, pmap, by_name, by_key)
        tiers[why].append((p, round(hours[p], 1)))
    P("## 對映預估（不寫）")
    for why in ("map", "exact", "key", "key+client", "ambiguous", "none", "bucket"):
        rows = tiers.get(why, [])
        P(f"- **{why}**：{len(rows)} 個，{round(sum(h for _, h in rows))} 小時")
    P("")
    P("## 撞案（owner 逐一指定 → PUT /timesheets/project_map）")
    for p, h in tiers.get("ambiguous", []):
        cands = by_key.get(sheet_project_key(p), [])
        P(f"- {p}（{h}h）→ 候選 {[(pid[:8], cli) for pid, cli in cands]}")
    P("")
    P("## 找不到（owner 決定：對既有案／建新案／當桶）—— 括號內是相似案名建議，不是對映")
    for p, h in tiers.get("none", []):
        sug = suggest_projects(p, by_key)
        P(f"- {p}（{h}h）" + (f" → 像：{sug}" if sug else ""))
    P("")
    staff = collections.Counter(r["staff"] for r in good)
    P(f"## 人員：{dict(staff)}（對不到 crm_staff 的在 ingest 回應 staff_unmatched）")


def _apply(api: Api, good: list, budgets: list, P) -> None:
    tot = collections.Counter()
    amb: set = set()
    unm: set = set()
    sun: set = set()
    for i in range(0, len(good), BATCH):
        resp = api.call("POST", "/api/v1/timesheets/ingest",
                        {"rows": good[i:i + BATCH], "source": "import"},
                        headers={"X-Timesheet-Token": api.token})
        for k in ("inserted", "skipped", "skipped_manual_priority"):
            tot[k] += resp.get(k, 0)
        amb |= set(resp.get("ambiguous_projects", []))
        unm |= set(resp.get("unmatched_projects", []))
        sun |= set(resp.get("staff_unmatched", []))
        print(f"  batch {i // BATCH + 1}: {resp.get('inserted')} 新增 / {resp.get('skipped')} 重複")
    P("## 匯入結果")
    P(f"- 新增 {tot['inserted']}、重複跳過 {tot['skipped']}、手填優先擋下 {tot['skipped_manual_priority']}")
    P(f"- 撞案 {len(amb)}：{sorted(amb)}")
    P(f"- 找不到 {len(unm)}：{sorted(unm)}")
    P(f"- 人員對不到 {len(sun)}：{sorted(sun)}")
    if budgets:
        r = api.call("PUT", "/api/v1/timesheets/budgets", {"items": budgets})
        P(f"- 預算：寫入 {r['applied']} 案；撞案 {len(r['ambiguous'])}；找不到 {len(r['unmatched'])}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--budget", action="store_true", help="一併從「專案狀態」灌 budget_hours")
    ap.add_argument("--apply", action="store_true", help="真的送；不帶＝只出報告（dry-run）")
    ap.add_argument("--prod", action="store_true", help="打生產 8000（預設 dev 8001）")
    ap.add_argument("--report", default=os.path.join(REPO, "docs", "timesheet_import_report.md"))
    a = ap.parse_args()

    good, bad = read_rows(a.xlsx)
    budgets = read_budgets(a.xlsx) if a.budget else []
    hours = collections.Counter()
    for r in good:
        hours[r["project"]] += r["hours"]

    out = io.StringIO()

    def P(*x):
        out.write(" ".join(str(y) for y in x) + "\n")

    mode = ("生產 8000" if a.prod else "dev 8001") + ("，APPLY" if a.apply else "，dry-run")
    P(f"# 工時匯入報告 {dt.datetime.now():%Y-%m-%d %H:%M}（{mode}）")
    P("")
    P(f"- 可送 {len(good)} 列 / {sum(hours.values()):.1f} 小時；壞列 {len(bad)}（不送）")
    P(f"- 人員：{dict(collections.Counter(r['staff'] for r in good).most_common())}")
    P(f"- 專案名 {len(hours)} 個；內部桶 {sorted(p for p in hours if p in INTERNAL_BUCKETS)}")
    if a.budget:
        P(f"- 預算：{len(budgets)} 案有值")
    P("")
    if bad:
        P("## 壞列（請到 Sheet 修）")
        for b in bad[:60]:
            P(f"- 列 {b['row']}：{b['date']!r} {b['staff']} {b['project']} hours={b['hours']!r}")
        P("")

    if a.apply:
        _apply(Api("http://127.0.0.1:8000" if a.prod else "http://127.0.0.1:8001"),
               good, budgets, P)
    else:
        _dry_run(good, hours, P, prod=a.prod)

    io.open(a.report, "w", encoding="utf-8").write(out.getvalue())
    print(out.getvalue())
    print(f"報告：{a.report}")


if __name__ == "__main__":
    main()
