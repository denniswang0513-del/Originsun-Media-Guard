"""core/doc_text.py — 從文件抽純文字（企劃範本消化用）。

只做一件事：檔案 → 乾淨的純文字。誰要用、抽完做什麼都不在這裡。

🔴 **康熙部首正規化不是可選項**。實測 owner 的真實範本（21 頁 PDF）：抽出來的
中文裡有 620 個字元是「康熙部首 / CJK 部首補充」——`⼭` 不是 `山`、`⿂` 不是
`魚`、`⼩` 不是 `小`。它們**長得一樣但不是同一個字元**，PDF 字型對映常吐這些。
不轉的話餵給 AI 的是半殘的文字，任何比對/搜尋也會失效。

`unicodedata.normalize("NFKC")` 修掉其中 608 個；剩下的 CJK Radicals
Supplement（U+2E80–U+2EFF）NFKC **不轉**（那個區塊沒有相容分解），要自己對映
——實測剩 `⻄`(U+2EC4)→`西`、`⺠`(U+2EA0)→`民`。這張表就是為了那些。

⚠️ NFKC 的副作用：全形「：」「／」會變成半形。對餵 AI 沒有差別，但如果哪天
要拿抽出來的文字去跟原檔逐字比對，要知道這件事。
"""
from __future__ import annotations

import os
import re
import unicodedata

# NFKC 轉不掉的 CJK 部首補充（U+2E80–U+2EFF）。實測到一個補一個 —— 硬編一張
# 「猜可能會遇到的」對照表只會讓人以為它是完整的。
_RADICAL_FIX = {
    "⻄": "西",   # CJK RADICAL WEST TWO
    "⺠": "民",   # CJK RADICAL CIVILIAN
}

SUPPORTED_EXTS = (".pdf", ".pptx", ".docx", ".xlsx", ".md", ".txt")


def normalize_cjk(text: str) -> str:
    """PDF/簡報抽出來的文字 → 正常的 CJK 字元（見模組檔頭）。"""
    out = unicodedata.normalize("NFKC", text or "")
    for bad, good in _RADICAL_FIX.items():
        out = out.replace(bad, good)
    # 抽出來常見的殘留：行尾空白、三行以上的空行
    out = re.sub(r"[ \t]+\n", "\n", out)
    return re.sub(r"\n{3,}", "\n\n", out).strip()


def _pdf(path: str) -> str:
    from pypdf import PdfReader
    return "\n".join((p.extract_text() or "") for p in PdfReader(path).pages)


def _pptx(path: str) -> str:
    """簡報：逐頁抽所有文字框 + 表格，頁與頁之間用分隔線。

    版面順序不等於閱讀順序（shapes 是 z-order），但對「抽出章節結構與語氣」
    這個用途夠用 —— 要精確閱讀順序得解版面座標，那是另一個量級的工程。
    """
    from pptx import Presentation
    out = []
    for i, slide in enumerate(Presentation(path).slides, 1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        if parts:
            out.append(f"--- 第 {i} 頁 ---\n" + "\n".join(parts))
    return "\n\n".join(out)


def _docx(path: str) -> str:
    import docx
    d = docx.Document(path)
    out = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out)


def _xlsx(path: str) -> str:
    """試算表：逐工作表逐列抽儲存格文字，格式比照 _pptx 的表格輸出
    （cell 以 ` | ` 相接、工作表之間標分隔線）。

    `data_only=True` —— 要的是算好的**值**不是公式字串（報價單全是 SUM），
    但代價是「只有 Excel 存檔時算過的值才在」：程式生成、從未用 Excel 開過的
    檔案，公式格會是 None → 被當空格跳過。read_only 串流模式：報價單常整張
    套格式，非 read_only 會把幾萬個空 styled cell 全實體化。
    """
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        out = []
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row
                         if c is not None and str(c).strip()]
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                out.append(f"--- 工作表 {ws.title} ---\n" + "\n".join(rows))
        return "\n\n".join(out)
    finally:
        wb.close()


def _plain(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


_READERS = {".pdf": _pdf, ".pptx": _pptx, ".docx": _docx, ".xlsx": _xlsx,
            ".md": _plain, ".txt": _plain}


def extract_text(path: str) -> tuple[str, str]:
    """檔案 → `(純文字, 錯誤訊息)`。成功時錯誤是 ""；失敗時文字是 ""。

    **同步**函式（要進 asyncio.to_thread）—— pypdf / python-pptx 都是 CPU-bound
    的純同步解析，21 頁的 PDF 約需數百毫秒，放在 event loop 上會卡住整台。
    """
    ext = os.path.splitext(path)[1].lower()
    reader = _READERS.get(ext)
    if not reader:
        if ext == ".xls":      # openpyxl 不吃舊二進位格式 —— 明講出路，別回通用訊息
            return "", ("舊版 Excel 格式 .xls 不支援（openpyxl 只讀 .xlsx）"
                        "—— 請用 Excel 另存成 .xlsx 再上傳")
        return "", f"不支援的格式 {ext}（只收 {'/'.join(SUPPORTED_EXTS)}）"
    try:
        raw = reader(path)
    except Exception as e:                       # 壞檔/加密 PDF/相依缺失都在這
        return "", f"{type(e).__name__}: {e}"
    text = normalize_cjk(raw)
    if not text:
        return "", "抽不到任何文字（可能整份是掃描圖片，需要先做 OCR）"
    return text, ""
