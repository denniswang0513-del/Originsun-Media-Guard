"""api_payroll.py — 人事管理 › 薪資：主檔／費率表／每月薪資單（docs/PAYROLL_OVERTIME_PLAN.md 第一批）。

權限：整支 check_admin_or_module(request, 'hr_payroll')（owner 2026-09-17「讓管理員勾選 預設合夥以上」）。
    薪資本來就是錢 —— 沒鑰匙一律 403，不走 CRM 的抹欄位那層（core/money.py `_PAYROLL_ONLY` 表態）。
帳本：只有母公司（entity='parent'）。薪資單「確認」時每列長一張請款單（crm_payment_requests，category 薪資／
    代發薪資、payee＝員工正式姓名＋身分證、payee_type 內部人員），之後就是既有的應付／匯款通知／銀行對帳。
NAS office-api 刻意不掛這支（main_office.py 的清單沒有它）。
"""
import io
import uuid
from datetime import datetime, timedelta, timezone

from fastapi import APIRouter, HTTPException, Request  # type: ignore
from fastapi.responses import StreamingResponse  # type: ignore
from sqlalchemy import func, select  # type: ignore

from core.auth import check_admin_or_module
from core.db_guard import db_factory_or_503
from core.hr_logic import active_staff_where
from core.payroll_logic import (DEFAULT_RATES, LINE_EDITABLE, PAY_TYPES, PAYROLL_ENTITIES, build_line, check_month, default_grades,
                                min_wage_warnings, next_month, normalize_rates, profile_as_of, vocab)
from core.schemas import PayProfileCreate, PayProfileUpdate, PayrollLineUpdate, PayrollRunCreate, RateTablePut
from db.models import CrmPaymentRequest, CrmStaff, PayrollLine, PayrollRateTable, PayrollRun, StaffPayProfile

router = APIRouter(prefix="/api/v1/hr/payroll", tags=["payroll"])

MODULE_KEY = "hr_payroll"
ENTITY = "parent"
PAYEE_TYPE = "內部人員"
CATEGORY_BY_ENTITY = {"公司": "薪資", "代發": "代發薪資"}   # 代發＝過帳不算費用（docs/CASHBOOK_CLASSIFICATION.md）
_ACTIVE_STAFF = active_staff_where()
_PROFILE_FIELDS = ("effective_from", "pay_type", "base_amount", "meal_allowance", "labor_grade", "health_grade",
                   "pension_self_rate", "dependents", "payroll_entity", "pay_day", "note")
_LINE_FIELDS = ("pay_type", "payroll_entity", "work_hours", "base_amount", "base_pay", "meal_allowance", "overtime_pay", "bonus_pay",
                "leave_deduction", "other_deduction", "labor_grade", "health_grade", "dependents", "pension_self_rate", "labor_self",
                "health_self", "pension_self", "labor_employer", "health_employer", "pension_employer", "gross_pay", "net_pay",
                "employer_total", "note")
_SUM_FIELDS = ("base_pay", "meal_allowance", "overtime_pay", "bonus_pay", "leave_deduction", "labor_self", "health_self",
               "pension_self", "other_deduction", "gross_pay", "net_pay", "labor_employer", "health_employer", "pension_employer",
               "employer_total")


def _guard(request: Request) -> dict:
    return check_admin_or_module(request, MODULE_KEY) or {}


def _actor(payload) -> str:
    return (payload or {}).get("sub") or ""


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _month_or_422(month: str) -> str:
    try:
        return check_month(month)
    except ValueError as e:
        raise HTTPException(status_code=422, detail=str(e))


def _profile_dict(p: StaffPayProfile) -> dict:
    return {"id": p.id, "staff_id": p.staff_id, "staff_name": p.staff_name or "", "effective_from": p.effective_from,
            "pay_type": p.pay_type or "月薪", "base_amount": int(p.base_amount or 0), "meal_allowance": int(p.meal_allowance or 0),
            "labor_grade": int(p.labor_grade or 0), "health_grade": int(p.health_grade or 0),
            "pension_self_rate": float(p.pension_self_rate or 0), "dependents": int(p.dependents or 0),
            "payroll_entity": p.payroll_entity or "公司", "pay_day": int(p.pay_day or 5), "note": p.note or ""}


def _line_dict(ln: PayrollLine) -> dict:
    d = {"id": ln.id, "run_id": ln.run_id, "staff_id": ln.staff_id, "staff_name": ln.staff_name or "", "profile_id": ln.profile_id or "",
         "manual_fields": list(ln.manual_fields or []), "bank_name": ln.bank_name or "", "bank_account": ln.bank_account or "",
         "payment_request_id": ln.payment_request_id or ""}
    for k in _LINE_FIELDS:
        v = getattr(ln, k)
        d[k] = (v or "") if k in ("pay_type", "payroll_entity", "note") else (float(v or 0) if k in ("work_hours", "pension_self_rate") else int(v or 0))
    return d


def _run_dict(r: PayrollRun, lines=None) -> dict:
    d = {"id": r.id, "month": r.month, "status": r.status or "草稿", "table_year": int(r.table_year or 0), "note": r.note or "",
         "confirmed_by": r.confirmed_by or "", "confirmed_at": r.confirmed_at.isoformat() if r.confirmed_at else None,
         "created_at": r.created_at.isoformat() if r.created_at else None}
    if lines is not None:
        rows = [_line_dict(x) for x in lines]
        d["lines"] = rows
        d["count"] = len(rows)
        d["totals"] = {k: sum(x[k] for x in rows) for k in _SUM_FIELDS}
    return d


async def _rates_for(session, year: int) -> dict:
    """該年的費率表；沒存過就給預設（DEFAULT_RATES 是 2026 版；別的年份也先用它，管理員在畫面改）。"""
    row = (await session.execute(select(PayrollRateTable).where(PayrollRateTable.year == year))).scalar_one_or_none()
    return normalize_rates(row.data if row else None, year)


def _check_profile_enums(pay_type: str, payroll_entity: str) -> None:
    if pay_type not in PAY_TYPES:
        raise HTTPException(status_code=422, detail=f"制度只能是 {'／'.join(PAY_TYPES)}")
    if payroll_entity not in PAYROLL_ENTITIES:
        raise HTTPException(status_code=422, detail=f"誰付只能是 {'／'.join(PAYROLL_ENTITIES)}")


# ── 薪資主檔 ──────────────────────────────────────────────────────────────

@router.get("/profiles")
async def list_profiles(request: Request, month: str = ""):
    """全員（在職）一列一人：`current`＝`month`（預設本月）適用那一段、`history`＝所有段。沒主檔的人也列（current=None）。"""
    _guard(request)
    month = _month_or_422(month) if month else datetime.now().strftime("%Y-%m")
    factory = db_factory_or_503()
    async with factory() as session:
        staff_rows = (await session.execute(select(CrmStaff).where(_ACTIVE_STAFF).order_by(CrmStaff.name))).scalars().all()
        profs = (await session.execute(select(StaffPayProfile).order_by(StaffPayProfile.staff_id, StaffPayProfile.effective_from))).scalars().all()
        by_staff: dict = {}
        for p in profs:
            by_staff.setdefault(p.staff_id, []).append(_profile_dict(p))
        rates = await _rates_for(session, int(month[:4]))
    out = []
    for s in staff_rows:
        hist = by_staff.get(s.id, [])
        out.append({"staff_id": s.id, "name": s.name, "role": s.role or "", "status": s.status or "在職",
                    "employment_type": s.employment_type or "", "has_bank": bool(s.bank_account),
                    "current": profile_as_of(hist, month), "history": hist})
    return {"month": month, "vocab": vocab(rates), "staff": out}


@router.post("/profiles")
async def create_profile(req: PayProfileCreate, request: Request):
    """新增一段主檔（一人一個 effective_from 只能有一段）；投保級距留空就照底薪＋伙食費帶。"""
    payload = _guard(request)
    eff = _month_or_422(req.effective_from)
    _check_profile_enums(req.pay_type, req.payroll_entity)
    factory = db_factory_or_503()
    async with factory() as session:
        staff = await session.get(CrmStaff, req.staff_id)
        if staff is None:
            raise HTTPException(status_code=404, detail="找不到這個人")
        dup = (await session.execute(select(StaffPayProfile.id).where(
            StaffPayProfile.staff_id == req.staff_id, StaffPayProfile.effective_from == eff))).first()
        if dup:
            raise HTTPException(status_code=409, detail=f"{staff.name} 從 {eff} 起已經有一段主檔，請改那一段")
        rates = await _rates_for(session, int(eff[:4]))
        grades = default_grades(req.pay_type, req.base_amount, req.meal_allowance, rates)
        p = StaffPayProfile(id=uuid.uuid4().hex, staff_id=staff.id, staff_name=staff.name, effective_from=eff,
                            pay_type=req.pay_type, base_amount=req.base_amount, meal_allowance=req.meal_allowance,
                            labor_grade=req.labor_grade if req.labor_grade is not None else grades["labor_grade"],
                            health_grade=req.health_grade if req.health_grade is not None else grades["health_grade"],
                            pension_self_rate=req.pension_self_rate, dependents=req.dependents,
                            payroll_entity=req.payroll_entity, pay_day=req.pay_day, note=req.note,
                            created_by=_actor(payload), created_at=_now(), updated_at=_now())
        session.add(p)
        await session.commit()
        return {"status": "ok", "profile": _profile_dict(p), "warnings": min_wage_warnings(p.pay_type, p.base_amount, rates)}


@router.put("/profiles/{profile_id}")
async def update_profile(profile_id: str, req: PayProfileUpdate, request: Request):
    """改一段主檔（只收給的欄位；staff_id 不能改）。回傳帶基本工資提醒。"""
    _guard(request)
    factory = db_factory_or_503()
    async with factory() as session:
        p = await session.get(StaffPayProfile, profile_id)
        if p is None:
            raise HTTPException(status_code=404, detail="找不到這段主檔")
        data = req.model_dump(exclude_none=True)
        if "effective_from" in data:
            data["effective_from"] = _month_or_422(data["effective_from"])
            dup = (await session.execute(select(StaffPayProfile.id).where(
                StaffPayProfile.staff_id == p.staff_id, StaffPayProfile.effective_from == data["effective_from"],
                StaffPayProfile.id != p.id))).first()
            if dup:
                raise HTTPException(status_code=409, detail="同一個月已經有另一段主檔")
        _check_profile_enums(data.get("pay_type", p.pay_type), data.get("payroll_entity", p.payroll_entity))
        for k in _PROFILE_FIELDS:
            if k in data:
                setattr(p, k, data[k])
        p.updated_at = _now()
        await session.commit()
        rates = await _rates_for(session, int(p.effective_from[:4]))
        return {"status": "ok", "profile": _profile_dict(p), "warnings": min_wage_warnings(p.pay_type, p.base_amount, rates)}


@router.delete("/profiles/{profile_id}")
async def delete_profile(profile_id: str, request: Request):
    """刪一段主檔。已產生的薪資單是快照，不受影響。"""
    _guard(request)
    factory = db_factory_or_503()
    async with factory() as session:
        p = await session.get(StaffPayProfile, profile_id)
        if p is None:
            raise HTTPException(status_code=404, detail="找不到這段主檔")
        await session.delete(p)
        await session.commit()
        return {"status": "ok"}


# ── 費率表 ────────────────────────────────────────────────────────────────

@router.get("/rates")
async def get_rates(request: Request, year: int = 0):
    """某年的費率表；沒存過就回系統預設（DEFAULT_RATES）並標 saved=False。"""
    _guard(request)
    year = year or datetime.now().year
    factory = db_factory_or_503()
    async with factory() as session:
        row = (await session.execute(select(PayrollRateTable).where(PayrollRateTable.year == year))).scalar_one_or_none()
        years = [y for (y,) in (await session.execute(select(PayrollRateTable.year).order_by(PayrollRateTable.year))).all()]
        return {"year": year, "saved": row is not None, "years": years, "default_year": DEFAULT_RATES["year"],
                "rates": normalize_rates(row.data if row else None, year)}


@router.put("/rates/{year}")
async def put_rates(year: int, req: RateTablePut, request: Request):
    """存某年的費率表（形狀由 normalize_rates 收斂：法定倍率與上限不給改）。"""
    payload = _guard(request)
    if not 2020 <= year <= 2100:
        raise HTTPException(status_code=422, detail="年份不對")
    data = normalize_rates(req.data, year)
    factory = db_factory_or_503()
    async with factory() as session:
        row = (await session.execute(select(PayrollRateTable).where(PayrollRateTable.year == year))).scalar_one_or_none()
        if row is None:
            row = PayrollRateTable(id=uuid.uuid4().hex, year=year, data=data, updated_by=_actor(payload), updated_at=_now())
            session.add(row)
        else:
            row.data, row.updated_by, row.updated_at = data, _actor(payload), _now()
        await session.commit()
        return {"status": "ok", "year": year, "rates": data}


# ── 每月薪資單 ────────────────────────────────────────────────────────────

async def _get_run(session, run_id: str, draft: bool = False) -> PayrollRun:
    r = await session.get(PayrollRun, run_id)
    if r is None:
        raise HTTPException(status_code=404, detail="找不到這張薪資單")
    if draft and r.status != "草稿":
        raise HTTPException(status_code=409, detail="已確認的薪資單不能改；要改請到請款單那邊處理")
    return r


async def _lines_of(session, run_id: str):
    return (await session.execute(select(PayrollLine).where(PayrollLine.run_id == run_id).order_by(PayrollLine.staff_name))).scalars().all()


async def _fill_lines(session, run: PayrollRun, keep: dict) -> list:
    """把主檔在 run.month 適用的人全部算成列。keep＝{staff_id: 既有列}：手改欄保留，其餘照主檔重算。回寫到的 staff_id。
    加班費：已核准、換加班費、掛這個月的加班單合計（routers.api_overtime.approved_pay_by_staff）；手改過 overtime_pay 的列不蓋。"""
    from routers.api_overtime import approved_pay_by_staff, mark_lines   # 懶 import：那邊核准時也會回頭呼叫這裡
    rates = await _rates_for(session, int(run.month[:4]))
    ot_pay = await approved_pay_by_staff(session, run.month)
    staff_rows = (await session.execute(select(CrmStaff).where(_ACTIVE_STAFF).order_by(CrmStaff.name))).scalars().all()
    profs = (await session.execute(select(StaffPayProfile))).scalars().all()
    by_staff: dict = {}
    for p in profs:
        by_staff.setdefault(p.staff_id, []).append(_profile_dict(p))
    written = []
    for s in staff_rows:
        prof = profile_as_of(by_staff.get(s.id, []), run.month)
        if not prof:
            continue
        old = keep.get(s.id)
        manual = list((old.manual_fields if old else None) or [])
        extras = {k: getattr(old, k) for k in manual} if old else {}
        if old and "work_hours" not in manual:
            extras["work_hours"] = old.work_hours       # 手填的時數不是「手改」，但也不該被重算洗掉
        if "overtime_pay" not in manual:
            extras["overtime_pay"] = ot_pay.get(s.id, 0)
        line = build_line(prof, rates, extras=extras)
        ln = old or PayrollLine(id=uuid.uuid4().hex, run_id=run.id, staff_id=s.id, created_at=_now())
        ln.staff_name, ln.profile_id = s.name, prof["id"]
        ln.bank_name, ln.bank_account = s.bank_name or "", s.bank_account or ""
        for k in _LINE_FIELDS:
            setattr(ln, k, line[k])
        ln.manual_fields = manual
        ln.updated_at = _now()
        if not old:
            session.add(ln)
        written.append(s.id)
    await session.flush()
    await mark_lines(session, run.month, {ln.staff_id: ln.id for ln in await _lines_of(session, run.id)})
    return written


async def refresh_staff_line(session, staff_id: str, month: str) -> bool:
    """加班核准後：`month` 的薪資單是草稿 → 只重算那個人那一列（手改欄照舊）。沒草稿或已確認 → False（下次產生時會帶）。"""
    run = (await session.execute(select(PayrollRun).where(PayrollRun.entity == ENTITY, PayrollRun.month == month))).scalar_one_or_none()
    if run is None or run.status != "草稿":
        return False
    old = {ln.staff_id: ln for ln in await _lines_of(session, run.id) if ln.staff_id == staff_id}
    rates = await _rates_for(session, int(month[:4]))
    from routers.api_overtime import approved_pay_by_staff, mark_lines
    ot_pay = await approved_pay_by_staff(session, month)
    ln = old.get(staff_id)
    if ln is None:
        return False
    prof = await session.get(StaffPayProfile, ln.profile_id) if ln.profile_id else None
    if prof is None:
        return False
    manual = list(ln.manual_fields or [])
    extras = {k: getattr(ln, k) for k in manual}
    extras["work_hours"] = ln.work_hours
    if "overtime_pay" not in manual:
        extras["overtime_pay"] = ot_pay.get(staff_id, 0)
    line = build_line(_profile_dict(prof), rates, extras=extras)
    for k in _LINE_FIELDS:
        setattr(ln, k, line[k])
    ln.updated_at = run.updated_at = _now()
    await mark_lines(session, month, {staff_id: ln.id})
    return True


@router.get("/runs")
async def list_runs(request: Request):
    """薪資單清單（新的月份在前）＋每張的人數、實發合計、公司總成本。"""
    _guard(request)
    factory = db_factory_or_503()
    async with factory() as session:
        runs = (await session.execute(select(PayrollRun).where(PayrollRun.entity == ENTITY).order_by(PayrollRun.month.desc()))).scalars().all()
        agg = {rid: (int(n), int(net or 0), int(emp or 0)) for rid, n, net, emp in (await session.execute(
            select(PayrollLine.run_id, func.count(PayrollLine.id), func.sum(PayrollLine.net_pay), func.sum(PayrollLine.employer_total))
            .group_by(PayrollLine.run_id))).all()}
        out = []
        for r in runs:
            d = _run_dict(r)
            n, net, emp = agg.get(r.id, (0, 0, 0))
            d.update({"count": n, "net_total": net, "employer_total": emp})
            out.append(d)
        return {"runs": out, "vocab": vocab(), "this_month": datetime.now().strftime("%Y-%m")}


@router.post("/runs")
async def create_run(req: PayrollRunCreate, request: Request):
    """產生某月薪資單草稿：主檔在該月適用的人各一列。同月已有 → 409。"""
    payload = _guard(request)
    month = _month_or_422(req.month)
    factory = db_factory_or_503()
    async with factory() as session:
        dup = (await session.execute(select(PayrollRun.id).where(PayrollRun.entity == ENTITY, PayrollRun.month == month))).first()
        if dup:
            raise HTTPException(status_code=409, detail=f"{month} 已經有薪資單了")
        run = PayrollRun(id=uuid.uuid4().hex, entity=ENTITY, month=month, status="草稿", table_year=int(month[:4]),
                         note=req.note, created_by=_actor(payload), created_at=_now(), updated_at=_now())
        session.add(run)
        written = await _fill_lines(session, run, {})
        if not written:
            raise HTTPException(status_code=422, detail="沒有任何人在這個月有薪資主檔，先到「薪資主檔」填")
        await session.commit()
        lines = await _lines_of(session, run.id)
        return {"status": "ok", "run": _run_dict(run, lines)}


@router.get("/runs/{run_id}")
async def get_run(run_id: str, request: Request):
    """一張薪資單的明細（每人一列）＋合計。"""
    _guard(request)
    factory = db_factory_or_503()
    async with factory() as session:
        run = await _get_run(session, run_id)
        lines = await _lines_of(session, run.id)
        return {"run": _run_dict(run, lines), "vocab": vocab(await _rates_for(session, int(run.month[:4])))}


@router.post("/runs/{run_id}/refresh")
async def refresh_run(run_id: str, request: Request):
    """草稿重算：主檔或費率改了之後按這顆；手改欄與手填時數保留，主檔已刪的人那列拿掉。"""
    _guard(request)
    factory = db_factory_or_503()
    async with factory() as session:
        run = await _get_run(session, run_id, draft=True)
        old = {ln.staff_id: ln for ln in await _lines_of(session, run.id)}
        written = set(await _fill_lines(session, run, old))
        for sid, ln in old.items():
            if sid not in written:
                await session.delete(ln)
        run.updated_at = _now()
        await session.commit()
        lines = await _lines_of(session, run.id)
        return {"status": "ok", "count": len(written), "run": _run_dict(run, lines)}


@router.put("/runs/{run_id}/lines/{line_id}")
async def update_line(run_id: str, line_id: str, req: PayrollLineUpdate, request: Request):
    """手填欄（時數／加班費／獎金／請假扣薪／其他扣款／備註）；改了就重算應發實發、記進 manual_fields。"""
    _guard(request)
    factory = db_factory_or_503()
    async with factory() as session:
        run = await _get_run(session, run_id, draft=True)
        ln = await session.get(PayrollLine, line_id)
        if ln is None or ln.run_id != run.id:
            raise HTTPException(status_code=404, detail="找不到這一列")
        data = {k: v for k, v in req.model_dump(exclude_none=True).items() if k in LINE_EDITABLE}
        if not data:
            raise HTTPException(status_code=422, detail="沒有可以改的欄位")
        rates = await _rates_for(session, int(run.month[:4]))
        prof = await session.get(StaffPayProfile, ln.profile_id) if ln.profile_id else None
        prof_d = _profile_dict(prof) if prof else {k: getattr(ln, k) for k in ("pay_type", "base_amount", "meal_allowance", "labor_grade",
                                                                                 "health_grade", "dependents", "pension_self_rate", "payroll_entity")}
        manual = set(ln.manual_fields or [])
        extras = {k: getattr(ln, k) for k in LINE_EDITABLE if k != "work_hours"}
        extras["work_hours"] = ln.work_hours
        extras.update(data)
        manual.update(k for k in data if k not in ("note", "work_hours"))
        line = build_line(prof_d, rates, extras=extras)
        for k in _LINE_FIELDS:
            setattr(ln, k, line[k])
        ln.manual_fields = sorted(manual)
        ln.updated_at = run.updated_at = _now()
        await session.commit()
        return {"status": "ok", "line": _line_dict(ln)}


@router.post("/runs/{run_id}/confirm")
async def confirm_run(run_id: str, request: Request):
    """草稿 → 已確認：每列長一張請款單（應付款、預計付款月＝下個月），之後走既有應付／匯款流程。"""
    payload = _guard(request)
    factory = db_factory_or_503()
    async with factory() as session:
        run = await _get_run(session, run_id, draft=True)
        lines = await _lines_of(session, run.id)
        if not lines:
            raise HTTPException(status_code=422, detail="這張薪資單沒有任何一列")
        staff = {s.id: s for s in (await session.execute(select(CrmStaff).where(CrmStaff.id.in_([ln.staff_id for ln in lines])))).scalars().all()}
        pay_month = next_month(run.month)
        # 權責：薪資是 run.month 的費用（request_date＝該月最後一天），錢是下個月付（planned_month）。月結鎖了就 409。
        expense_day = datetime(int(pay_month[:4]), int(pay_month[5:]), 1, tzinfo=timezone.utc) - timedelta(days=1)
        from routers.crm._shared import _assert_month_open
        await _assert_month_open(session, expense_day, entity=ENTITY)
        now = _now()
        made = 0
        for ln in lines:
            if ln.payment_request_id or ln.net_pay <= 0:
                continue
            s = staff.get(ln.staff_id)
            pr = CrmPaymentRequest(
                id=uuid.uuid4().hex, entity=ENTITY,
                request_date=expense_day,
                amount=int(ln.net_pay), summary=f"{run.month} 薪資 {ln.staff_name}",
                category=CATEGORY_BY_ENTITY.get(ln.payroll_entity or "公司", "薪資"),
                payee_name=ln.staff_name, payee_id=(s.id_number if s else "") or "", payee_type=PAYEE_TYPE,
                payment_status="應付款", planned_month=pay_month, created_at=now, updated_at=now)
            session.add(pr)
            ln.payment_request_id = pr.id
            ln.updated_at = now
            made += 1
        run.status, run.confirmed_by, run.confirmed_at, run.updated_at = "已確認", _actor(payload), now, now
        await session.commit()
        lines = await _lines_of(session, run.id)
        return {"status": "ok", "payment_requests": made, "run": _run_dict(run, lines)}


@router.delete("/runs/{run_id}")
async def delete_run(run_id: str, request: Request):
    """刪掉草稿（已確認的不給刪 —— 請款單已經長出去了）。"""
    _guard(request)
    factory = db_factory_or_503()
    async with factory() as session:
        run = await _get_run(session, run_id, draft=True)
        for ln in await _lines_of(session, run.id):
            await session.delete(ln)
        await session.delete(run)
        await session.commit()
        return {"status": "ok"}


_EXPORT_COLS = (("staff_name", "姓名"), ("pay_type", "制度"), ("work_hours", "時數"), ("base_pay", "底薪"), ("meal_allowance", "伙食費"),
                ("overtime_pay", "加班費"), ("bonus_pay", "獎金"), ("leave_deduction", "請假扣薪"), ("gross_pay", "應發"),
                ("labor_self", "勞保自負"), ("health_self", "健保自負"), ("pension_self", "勞退自提"), ("other_deduction", "其他扣款"),
                ("net_pay", "實發"), ("labor_employer", "雇主勞保"), ("health_employer", "雇主健保"), ("pension_employer", "雇主勞退"),
                ("employer_total", "公司總成本"), ("payroll_entity", "誰付"), ("bank_name", "銀行"), ("bank_account", "帳號"), ("note", "備註"))


@router.get("/runs/{run_id}/export")
async def export_run(run_id: str, request: Request):
    """印領清冊（Excel）給會計：一列一人，最後一列合計。"""
    _guard(request)
    from openpyxl import Workbook  # type: ignore
    from openpyxl.styles import Font  # type: ignore
    factory = db_factory_or_503()
    async with factory() as session:
        run = await _get_run(session, run_id)
        lines = [_line_dict(x) for x in await _lines_of(session, run.id)]
    wb = Workbook()
    ws = wb.active
    ws.title = f"{run.month} 薪資清冊"
    ws.append([f"源日有限公司 {run.month} 薪資印領清冊（{run.status}）"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([label for _k, label in _EXPORT_COLS])
    for c in ws[2]:
        c.font = Font(bold=True)
    for d in lines:
        ws.append([d.get(k, "") for k, _l in _EXPORT_COLS])
    total = ["合計", "", ""] + [sum(d[k] for d in lines) if k in _SUM_FIELDS else "" for k, _l in _EXPORT_COLS[3:]]
    ws.append(total)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"payroll_{run.month}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}"'})
